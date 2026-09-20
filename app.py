import html
import os
import re
import shutil
import subprocess
import textwrap
from pathlib import Path

import streamlit as st
from docx import Document
from reportlab.lib.pagesizes import LETTER
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer

from utils.pdf_reader import extract_text
from utils.ai import analyze_resume
from utils.jd_matcher import match_resume_to_job
from utils.resume_store import build_resume_store
from utils.vector_store import load_vectors
from rag import answer_question


# ============================================================
# APP CONSTANTS
# ============================================================

APP_NAME = "AI Resume Reviewer"
RESUME_ID = "sivani_resume"
TEMP_RESUME_PATH = "temp_resume.pdf"
TEMP_INDEX_PDF_PATH = "temp_resume_for_indexing.pdf"

GITHUB_URL = "https://github.com/sivanisankar123/AI-Resume-Reviewer"
LIVE_URL = "https://ai-resume-reviewer-5wxpt59pgrevmtfskknhhl.streamlit.app/"
LINKEDIN_URL = "https://www.linkedin.com/in/sivani-sankar-mohapatra-9879a1135"


# ============================================================
# PAGE CONFIG
# ============================================================

st.set_page_config(
    page_title=APP_NAME,
    page_icon="📄",
    layout="wide",
    initial_sidebar_state="expanded",
)


# ============================================================
# HELPERS
# ============================================================

def esc(value) -> str:
    return html.escape(str(value))


def safe_list(value) -> list:
    return value if isinstance(value, list) else []


def render_section_title(icon: str, title: str, subtitle: str = ""):
    st.markdown(
        f"""
        <div class="section-heading">
            <div class="section-icon">{icon}</div>
            <div>
                <div class="section-title">{esc(title)}</div>
                <div class="section-subtitle">{esc(subtitle)}</div>
            </div>
        </div>
        """,
        unsafe_allow_html=True,
    )


def render_metric_card(icon: str, label: str, value: str, tone: str = "blue"):
    st.markdown(
        f"""
        <div class="metric-card tone-{tone}">
            <div class="metric-top">
                <div class="metric-icon">{icon}</div>
                <div class="metric-label">{esc(label)}</div>
            </div>
            <div class="metric-value">{esc(value)}</div>
        </div>
        """,
        unsafe_allow_html=True,
    )


def render_list_card(title: str, icon: str, items: list, empty_text: str = "No items available."):
    st.markdown(
        f"""
        <div class="list-card-title">{icon} {esc(title)}</div>
        """,
        unsafe_allow_html=True,
    )

    if not items:
        st.caption(empty_text)
        return

    for item in items:
        st.markdown(
            f"""
            <div class="list-row">
                <span class="bullet-dot"></span>
                <span>{esc(item)}</span>
            </div>
            """,
            unsafe_allow_html=True,
        )


def normalize_resume_id(filename: str) -> str:
    name = filename.rsplit(".", 1)[0].lower().strip()
    name = re.sub(r"[^a-z0-9]+", "_", name)
    name = name.strip("_")
    return name or RESUME_ID


def extract_docx_text(path: str) -> str:
    """Extract readable text from a DOCX resume, including table content."""
    document = Document(path)
    parts = []

    for paragraph in document.paragraphs:
        value = paragraph.text.strip()
        if value:
            parts.append(value)

    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            row_text = " | ".join(cell for cell in cells if cell)
            if row_text:
                parts.append(row_text)

    return "\\n".join(parts).strip()


def extract_doc_text(path: str) -> str:
    """Best-effort extraction for legacy .doc files.

    Streamlit Cloud does not guarantee a Microsoft Word runtime, so try
    common command-line converters when available and fail with a clear
    message instead of silently returning bad text.
    """
    for command in ("antiword", "catdoc"):
        executable = shutil.which(command)
        if executable:
            try:
                result = subprocess.run(
                    [executable, path],
                    capture_output=True,
                    text=True,
                    timeout=30,
                    check=True,
                )
                if result.stdout.strip():
                    return result.stdout.strip()
            except (subprocess.SubprocessError, OSError):
                pass

    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if soffice:
        out_dir = str(Path(path).parent / "doc_conversion")
        os.makedirs(out_dir, exist_ok=True)
        try:
            subprocess.run(
                [soffice, "--headless", "--convert-to", "txt:Text", "--outdir", out_dir, path],
                capture_output=True,
                text=True,
                timeout=60,
                check=True,
            )
            converted = Path(out_dir) / (Path(path).stem + ".txt")
            if converted.exists():
                value = converted.read_text(errors="ignore").strip()
                if value:
                    return value
        except (subprocess.SubprocessError, OSError):
            pass

    raise RuntimeError(
        "Legacy .doc extraction is unavailable in this deployment environment. "
        "Please use .docx or .pdf, or install antiword/LibreOffice for .doc support."
    )


def write_text_as_pdf(text: str, output_path: str) -> None:
    """Create a simple PDF from extracted DOC/DOCX text for the existing RAG pipeline."""
    styles = getSampleStyleSheet()
    body = styles["BodyText"]
    body.fontName = "Helvetica"
    body.fontSize = 9
    body.leading = 12

    document = SimpleDocTemplate(
        output_path,
        pagesize=LETTER,
        rightMargin=40,
        leftMargin=40,
        topMargin=40,
        bottomMargin=40,
    )

    story = []
    for paragraph in text.splitlines():
        cleaned = paragraph.strip()
        if cleaned:
            safe = html.escape(cleaned).replace("\\n", "<br/>")
            story.append(Paragraph(safe, body))
            story.append(Spacer(1, 4))

    document.build(story)


def prepare_uploaded_resume(uploaded_file) -> str:
    """Save the upload, extract its text, and prepare a PDF for the existing RAG builder."""
    suffix = Path(uploaded_file.name).suffix.lower()
    upload_path = f"temp_resume{suffix}"

    with open(upload_path, "wb") as file:
        file.write(uploaded_file.getbuffer())

    if suffix == ".pdf":
        # The upload is already saved as temp_resume.pdf. Do not copy the file
        # onto itself, which raises SameFileError.
        if upload_path != TEMP_RESUME_PATH:
            shutil.copyfile(upload_path, TEMP_RESUME_PATH)
        return extract_text(TEMP_RESUME_PATH)

    if suffix == ".docx":
        resume_text = extract_docx_text(upload_path)
    elif suffix == ".doc":
        resume_text = extract_doc_text(upload_path)
    else:
        raise ValueError("Unsupported resume format. Please upload PDF, DOC or DOCX.")

    if not resume_text.strip():
        raise ValueError("No readable text was found in the uploaded resume.")

    # Convert extracted text to a temporary PDF so the existing vector-store
    # pipeline can continue using its current PDF-based implementation.
    write_text_as_pdf(resume_text, TEMP_INDEX_PDF_PATH)
    shutil.copyfile(TEMP_INDEX_PDF_PATH, TEMP_RESUME_PATH)
    return resume_text


def build_current_resume_store() -> int:
    """Build the existing vector store from the normalized temporary PDF."""
    return build_resume_store(TEMP_RESUME_PATH, st.session_state.resume_id)


def render_score_ring(score: int):
    score = max(0, min(100, int(score)))
    st.markdown(
        f"""
        <div class="score-wrap">
            <div class="score-ring" style="--score:{score * 3.6}deg;">
                <div class="score-center">
                    <div class="score-number">{score}%</div>
                    <div class="score-caption">ATS Score</div>
                </div>
            </div>
        </div>
        """,
        unsafe_allow_html=True,
    )


# ============================================================
# PREMIUM DASHBOARD STYLES
# ============================================================

st.markdown(
    """
<style>
/* ---------- Base ---------- */
:root {
    --navy: #0f1b33;
    --navy-2: #162745;
    --blue: #2f6df6;
    --blue-2: #4d83ff;
    --indigo: #6756f5;
    --text: #17213a;
    --muted: #68758f;
    --surface: #ffffff;
    --surface-2: #f6f9ff;
    --border: #e2e9f4;
    --green: #11a36a;
    --purple: #8b5cf6;
    --amber: #f0a51a;
    --red: #e95555;
}

.stApp {
    background:
        radial-gradient(circle at 90% 0%, rgba(72, 116, 255, 0.12), transparent 24rem),
        linear-gradient(180deg, #f4f7fd 0%, #fbfcff 44%, #f8faff 100%);
    color: var(--text);
}

.block-container {
    max-width: 1480px;
    padding-top: 1.2rem;
    padding-bottom: 1.2rem;
}

/* ---------- Hide Streamlit chrome ---------- */
#MainMenu {visibility: hidden;}
footer {visibility: hidden;}
header[data-testid="stHeader"] {
    background: transparent;
}

/* ---------- Sidebar ---------- */
section[data-testid="stSidebar"] {
    background:
        linear-gradient(180deg, #0f1b33 0%, #122343 58%, #10203c 100%);
    border-right: 1px solid rgba(255,255,255,0.06);
}

section[data-testid="stSidebar"] > div {
    padding-top: 1.3rem;
}

section[data-testid="stSidebar"] * {
    color: #f7f9ff;
}

.sidebar-brand {
    padding: 4px 6px 18px 6px;
}

.sidebar-brand-title {
    font-size: 1.32rem;
    font-weight: 800;
    letter-spacing: -0.02em;
}

.sidebar-brand-subtitle {
    color: #a9b9d6;
    margin-top: 5px;
    font-size: .88rem;
}

.sidebar-tagline {
    color: #7fa2df;
    font-size: .75rem;
    font-style: italic;
    margin-top: 8px;
}

.sidebar-divider {
    height: 1px;
    background: rgba(255,255,255,0.08);
    margin: 14px 0;
}

.nav-title {
    color: #8296b8;
    font-size: .72rem;
    text-transform: uppercase;
    letter-spacing: .08em;
    font-weight: 800;
    margin: 8px 0 10px;
}

.nav-link {
    display: flex;
    align-items: center;
    gap: 11px;
    padding: 10px 12px;
    border-radius: 10px;
    color: #dce6f7 !important;
    text-decoration: none !important;
    margin: 4px 0;
    font-size: .91rem;
    transition: all .18s ease;
}

.nav-link:hover {
    background: rgba(83, 133, 255, .17);
    color: white !important;
}

.nav-link-active {
    background: linear-gradient(90deg, #2e6df6, #4384ff);
    box-shadow: 0 8px 18px rgba(47,109,246,.23);
    color: white !important;
}

.sidebar-tip {
    display: flex;
    gap: 9px;
    color: #c7d4e8;
    font-size: .82rem;
    line-height: 1.42;
    margin: 9px 0;
}

.sidebar-quote {
    border: 1px solid rgba(130, 160, 220, .25);
    background: rgba(255,255,255,.035);
    padding: 14px;
    border-radius: 12px;
    color: #c7d6ed;
    font-size: .85rem;
    font-style: italic;
    margin-top: 18px;
}

/* ---------- Hero ---------- */
.hero {
    min-height: 180px;
    border-radius: 24px;
    overflow: hidden;
    padding: 30px 34px;
    margin-bottom: 18px;
    position: relative;
    background:
        radial-gradient(circle at 84% 20%, rgba(104, 96, 255, .22), transparent 16rem),
        radial-gradient(circle at 65% 100%, rgba(73, 133, 255, .18), transparent 18rem),
        linear-gradient(135deg, #ffffff 0%, #f3f7ff 46%, #edf3ff 100%);
    border: 1px solid #dfe8f7;
    box-shadow: 0 12px 34px rgba(32, 59, 112, .08);
}

.hero::after {
    content: "";
    position: absolute;
    width: 320px;
    height: 320px;
    right: -65px;
    top: -95px;
    border-radius: 50%;
    border: 42px solid rgba(76, 121, 255, .07);
}

.hero-kicker {
    color: #6b7a95;
    font-size: .82rem;
    font-weight: 750;
    text-transform: uppercase;
    letter-spacing: .08em;
    margin-bottom: 7px;
}

.hero-title {
    font-size: 2.25rem;
    line-height: 1.05;
    font-weight: 850;
    letter-spacing: -.045em;
    color: #10204b;
}

.hero-title span {
    color: #2f6df6;
}

.hero-subtitle {
    font-size: 1rem;
    color: #5e6d89;
    max-width: 760px;
    margin-top: 10px;
    line-height: 1.6;
}

.hero-badge {
    display: inline-flex;
    align-items: center;
    gap: 7px;
    margin-top: 17px;
    padding: 7px 11px;
    border-radius: 999px;
    background: white;
    border: 1px solid #dae4f4;
    color: #2a4c9a;
    font-weight: 700;
    font-size: .78rem;
}

/* ---------- Feature strip ---------- */
.feature-strip {
    display: grid;
    grid-template-columns: repeat(4, minmax(0, 1fr));
    gap: 12px;
    margin: 12px 0 19px;
}

.feature-chip {
    display: flex;
    align-items: center;
    gap: 12px;
    background: rgba(255,255,255,.86);
    border: 1px solid #e2e9f4;
    border-radius: 15px;
    padding: 12px 14px;
    box-shadow: 0 6px 18px rgba(30, 48, 94, .045);
}

.feature-chip-icon {
    width: 43px;
    height: 43px;
    display: grid;
    place-items: center;
    border-radius: 12px;
    font-size: 1.25rem;
    background: #edf3ff;
}

.feature-chip-title {
    font-weight: 800;
    color: #1c2a50;
    font-size: .92rem;
}

.feature-chip-subtitle {
    color: #8290a8;
    font-size: .72rem;
    margin-top: 2px;
}

/* ---------- Section headings ---------- */
.section-heading {
    display: flex;
    align-items: center;
    gap: 11px;
    margin: 7px 0 13px;
}

.section-icon {
    width: 42px;
    height: 42px;
    display: grid;
    place-items: center;
    border-radius: 12px;
    background: #eef3ff;
    font-size: 1.2rem;
}

.section-title {
    font-weight: 850;
    font-size: 1.15rem;
    color: #14224a;
    letter-spacing: -.02em;
}

.section-subtitle {
    color: #7e8ca4;
    font-size: .78rem;
    margin-top: 2px;
}

/* ---------- Streamlit bordered containers as cards ---------- */
[data-testid="stVerticalBlockBorderWrapper"] {
    background: rgba(255,255,255,.92);
    border: 1px solid #e0e8f4 !important;
    border-radius: 18px !important;
    box-shadow: 0 9px 24px rgba(24, 50, 93, .055);
}

[data-testid="stVerticalBlockBorderWrapper"] > div {
    padding: 4px 5px 5px;
}

/* ---------- File uploader ---------- */
[data-testid="stFileUploader"] {
    border-radius: 14px;
}

[data-testid="stFileUploaderDropzone"] {
    border: 1.5px dashed #9dbbf8;
    border-radius: 14px;
    background: linear-gradient(180deg, #fbfdff, #f6f9ff);
    min-height: 180px;
}

[data-testid="stFileUploaderDropzone"] button {
    background: linear-gradient(135deg, #2f6df6, #4e83ff);
    color: white;
    border: 0;
    border-radius: 9px;
}

/* ---------- Buttons ---------- */
div.stButton > button,
div.stDownloadButton > button {
    min-height: 42px;
    border-radius: 10px;
    font-weight: 750;
    border: 1px solid #dce5f3;
    transition: all .15s ease;
}

div.stButton > button:hover,
div.stDownloadButton > button:hover {
    transform: translateY(-1px);
    border-color: #8eb0fa;
    box-shadow: 0 7px 18px rgba(54,100,211,.11);
}

/* ---------- Metrics ---------- */
.metric-card {
    border: 1px solid #e1e9f5;
    background: linear-gradient(180deg, #ffffff, #fbfcff);
    border-radius: 15px;
    padding: 14px 15px;
    min-height: 112px;
    box-shadow: 0 7px 18px rgba(40, 60, 100, .04);
}

.metric-top {
    display: flex;
    align-items: center;
    gap: 8px;
}

.metric-icon {
    width: 34px;
    height: 34px;
    border-radius: 9px;
    display: grid;
    place-items: center;
    background: #eef4ff;
}

.metric-label {
    font-size: .76rem;
    color: #7b8aa4;
    font-weight: 720;
}

.metric-value {
    color: #14224a;
    font-size: 1.65rem;
    font-weight: 850;
    margin-top: 10px;
}

.tone-green .metric-icon {background:#e9fbf3;}
.tone-purple .metric-icon {background:#f3edff;}
.tone-amber .metric-icon {background:#fff7e5;}
.tone-red .metric-icon {background:#fff0f0;}

/* ---------- Score ring ---------- */
.score-wrap {
    display: flex;
    align-items: center;
    justify-content: center;
    padding: 8px 0;
}

.score-ring {
    width: 125px;
    height: 125px;
    border-radius: 50%;
    background: conic-gradient(#17a669 var(--score), #e9eef6 0);
    display: grid;
    place-items: center;
    box-shadow: 0 8px 20px rgba(22,166,105,.10);
}

.score-ring::before {
    content: "";
    position: absolute;
    width: 93px;
    height: 93px;
    border-radius: 50%;
    background: white;
}

.score-center {
    z-index: 2;
    text-align: center;
}

.score-number {
    color: #14224a;
    font-weight: 850;
    font-size: 1.5rem;
}

.score-caption {
    color: #7d899f;
    font-size: .72rem;
}

/* ---------- Result list cards ---------- */
.list-card-title {
    color: #17264e;
    font-weight: 820;
    margin: 2px 0 10px;
}

.list-row {
    display: flex;
    gap: 9px;
    align-items: flex-start;
    padding: 7px 0;
    border-bottom: 1px solid #eef2f8;
    color: #44536e;
    line-height: 1.45;
    font-size: .88rem;
}

.list-row:last-child {
    border-bottom: 0;
}

.bullet-dot {
    min-width: 7px;
    min-height: 7px;
    border-radius: 50%;
    background: #4c7ff4;
    margin-top: 7px;
}

/* ---------- Quick insights ---------- */
.insight-row {
    display:flex;
    gap:10px;
    align-items:flex-start;
    margin: 12px 0;
}

.insight-icon {
    width:35px;
    height:35px;
    border-radius:10px;
    display:grid;
    place-items:center;
    background:#eef4ff;
}

.insight-title {
    font-weight:800;
    color:#21325f;
    font-size:.86rem;
}

.insight-text {
    color:#8390a8;
    font-size:.75rem;
}

/* ---------- Knowledge state ---------- */
.kb-ready {
    padding: 18px;
    border-radius: 14px;
    border: 1px solid #cfeee0;
    background: linear-gradient(135deg, #effbf5, #f8fffb);
}

.kb-ready-title {
    font-weight: 850;
    color: #08774d;
    margin-bottom: 4px;
}

.kb-ready-text {
    color: #527566;
    font-size: .82rem;
}

/* ---------- Chat ---------- */
[data-testid="stChatMessage"] {
    background: rgba(255,255,255,.92);
    border: 1px solid #e2e9f4;
    border-radius: 15px;
    padding: 8px 10px;
    margin: 7px 0;
}

[data-testid="stChatInput"] {
    border-radius: 13px;
}

/* ---------- Footer ---------- */
.app-footer {
    border-top: 1px solid #e3e9f3;
    margin-top: 28px;
    padding: 18px 4px 6px;
    display: grid;
    grid-template-columns: 1.2fr .9fr .9fr;
    gap: 18px;
    align-items: center;
    color: #7e8aa0;
    font-size: .76rem;
}

.footer-name {
    text-align: right;
}

.footer-name strong {
    color: #1a3475;
    display: block;
    font-size: .84rem;
}

.footer-center {
    text-align:center;
    color:#5972a2;
    font-style:italic;
}

.footer-links a {
    color: #315fbd !important;
    text-decoration: none;
    font-weight: 700;
    margin-right: 12px;
}

/* ---------- Strong button contrast ---------- */
div.stButton > button,
div.stButton > button[kind="secondary"],
div.stButton > button[kind="primary"] {
    background: linear-gradient(135deg, #2f6df6, #4b82ff) !important;
    color: #ffffff !important;
    -webkit-text-fill-color: #ffffff !important;
    border: 1px solid #2f6df6 !important;
    font-weight: 800 !important;
    box-shadow: 0 5px 14px rgba(47,109,246,.16) !important;
}

div.stButton > button:hover,
div.stButton > button[kind="secondary"]:hover,
div.stButton > button[kind="primary"]:hover {
    background: linear-gradient(135deg, #245bd1, #3b72ea) !important;
    color: #ffffff !important;
    -webkit-text-fill-color: #ffffff !important;
}

/* ---------- Readable text areas on light and dark Streamlit themes ---------- */
[data-testid="stTextArea"] textarea {
    background: #ffffff !important;
    color: #17213a !important;
    -webkit-text-fill-color: #17213a !important;
    caret-color: #2f6df6 !important;
    border: 1px solid #dce5f3 !important;
    border-radius: 12px !important;
    line-height: 1.55 !important;
}

[data-testid="stTextArea"] textarea::placeholder {
    color: #8a97ac !important;
    -webkit-text-fill-color: #8a97ac !important;
}

[data-testid="stTextArea"] textarea:focus {
    border-color: #7ca4f8 !important;
    box-shadow: 0 0 0 2px rgba(47,109,246,.10) !important;
}

/* ---------- Guided workflow ---------- */
.workflow-card {
    background: rgba(255,255,255,.94);
    border: 1px solid #dfe7f4;
    border-radius: 18px;
    padding: 20px;
    margin: 0 0 20px;
    box-shadow: 0 8px 22px rgba(30, 48, 94, .05);
}

.workflow-title {
    color: #152650;
    font-size: 1.05rem;
    font-weight: 850;
    margin-bottom: 4px;
}

.workflow-subtitle {
    color: #75839d;
    font-size: .8rem;
    margin-bottom: 15px;
}

.workflow-grid {
    display: grid;
    grid-template-columns: repeat(5, minmax(0, 1fr));
    gap: 10px;
}

.workflow-step {
    position: relative;
    background: #f7f9fe;
    border: 1px solid #e3eaf6;
    border-radius: 13px;
    padding: 13px;
    min-height: 118px;
}

.workflow-number {
    width: 27px;
    height: 27px;
    border-radius: 50%;
    display: grid;
    place-items: center;
    background: #2f6df6;
    color: #fff;
    font-size: .75rem;
    font-weight: 850;
    margin-bottom: 8px;
}

.workflow-step-title {
    color: #1a2b55;
    font-weight: 820;
    font-size: .84rem;
}

.workflow-step-text {
    color: #7a879f;
    font-size: .72rem;
    line-height: 1.45;
    margin-top: 4px;
}

.workflow-next {
    color: #7290c9;
    font-size: .72rem;
    margin-top: 7px;
    font-weight: 750;
}

/* ---------- File support badge ---------- */
.file-support {
    display: inline-flex;
    gap: 7px;
    flex-wrap: wrap;
    margin: 3px 0 12px;
}

.file-type {
    background: #eef4ff;
    color: #315ba8;
    border: 1px solid #d8e5ff;
    border-radius: 999px;
    padding: 4px 9px;
    font-size: .7rem;
    font-weight: 800;
}

/* ---------- Responsive ---------- */
@media (max-width: 1100px) {
    .workflow-grid {
        grid-template-columns: repeat(3, minmax(0, 1fr));
    }
}

@media (max-width: 900px) {
    .workflow-grid {
        grid-template-columns: 1fr;
    }
    .workflow-step {
        min-height: auto;
    }
    .feature-strip {
        grid-template-columns: repeat(2, minmax(0, 1fr));
    }
    .hero-title {
        font-size: 1.8rem;
    }
    .app-footer {
        grid-template-columns: 1fr;
        text-align:left;
    }
    .footer-center, .footer-name {
        text-align:left;
    }
}
</style>
""",
    unsafe_allow_html=True,
)


# ============================================================
# SESSION STATE
# ============================================================

defaults = {
    "uploaded_file_name": None,
    "resume_text": "",
    "resume_id": RESUME_ID,
    "analysis": None,
    "job_result": None,
    "knowledge_base_ready": False,
    "knowledge_base_chunks": 0,
    "resume_conversation": [],
    "job_description": "",
}

for key, value in defaults.items():
    if key not in st.session_state:
        st.session_state[key] = value


# ============================================================
# SIDEBAR
# ============================================================

with st.sidebar:
    st.markdown(
        """
        <div class="sidebar-brand">
            <div class="sidebar-brand-title">📄 AI Resume Reviewer</div>
            <div class="sidebar-brand-subtitle">GenAI • RAG • Agentic AI</div>
            <div class="sidebar-tagline">Smarter Resumes. Brighter Opportunities.</div>
        </div>
        <div class="sidebar-divider"></div>

        <div class="nav-title">Navigation</div>
        <a class="nav-link nav-link-active" href="#home">🏠 <span>Home</span></a>
        <a class="nav-link" href="#resume-analysis">📄 <span>Resume Analysis</span></a>
        <a class="nav-link" href="#job-description-match">🎯 <span>Job Description Match</span></a>
        <a class="nav-link" href="#resume-knowledge-base">🧠 <span>Resume Knowledge Base</span></a>
        <a class="nav-link" href="#resume-agent">💬 <span>Ask Resume Agent</span></a>
        <a class="nav-link" href="#about">ℹ️ <span>About</span></a>

        <div class="sidebar-divider"></div>
        <div class="nav-title">Quick Tips</div>

        <div class="sidebar-tip">💡 <span>Upload a clear, updated PDF resume.</span></div>
        <div class="sidebar-tip">🎯 <span>Use a specific job description for better insights.</span></div>
        <div class="sidebar-tip">🤖 <span>Ask the Resume Agent anything about your resume.</span></div>

        <div class="sidebar-quote">
            “AI turns information into opportunities.”
        </div>
        """,
        unsafe_allow_html=True,
    )


# ============================================================
# HOME / HERO
# ============================================================

st.markdown('<div id="home"></div>', unsafe_allow_html=True)

st.markdown(
    """
    <div class="hero">
        <div class="hero-kicker">AI-powered career intelligence</div>
        <div class="hero-title">Welcome to <span>AI Resume Reviewer</span></div>
        <div class="hero-subtitle">
            Get AI-powered insights, ATS scoring, skill analysis, job-description matching,
            semantic resume search and personalized recommendations — all in one workspace.
        </div>
        <div class="hero-badge">✨ Turn your resume into opportunities</div>
    </div>

    <div class="feature-strip">
        <div class="feature-chip">
            <div class="feature-chip-icon">📄</div>
            <div><div class="feature-chip-title">Analyze</div><div class="feature-chip-subtitle">Upload your resume</div></div>
        </div>
        <div class="feature-chip">
            <div class="feature-chip-icon">🎯</div>
            <div><div class="feature-chip-title">Match</div><div class="feature-chip-subtitle">Check JD compatibility</div></div>
        </div>
        <div class="feature-chip">
            <div class="feature-chip-icon">💡</div>
            <div><div class="feature-chip-title">Improve</div><div class="feature-chip-subtitle">Get actionable suggestions</div></div>
        </div>
        <div class="feature-chip">
            <div class="feature-chip-icon">💬</div>
            <div><div class="feature-chip-title">Ask</div><div class="feature-chip-subtitle">Chat with your resume</div></div>
        </div>
    </div>

    <div class="workflow-card">
        <div class="workflow-title">🚀 How to use AI Resume Reviewer</div>
        <div class="workflow-subtitle">
            Follow these steps in order. Job matching is optional, while the Knowledge Base is required for Resume Agent chat.
        </div>
        <div class="workflow-grid">
            <div class="workflow-step">
                <div class="workflow-number">1</div>
                <div class="workflow-step-title">📄 Upload Resume</div>
                <div class="workflow-step-text">Upload your PDF, DOC or DOCX resume.</div>
                <div class="workflow-next">Start here →</div>
            </div>
            <div class="workflow-step">
                <div class="workflow-number">2</div>
                <div class="workflow-step-title">✨ Analyze Resume</div>
                <div class="workflow-step-text">Click Analyze Resume to get ATS score, experience, strengths and gaps.</div>
                <div class="workflow-next">Next →</div>
            </div>
            <div class="workflow-step">
                <div class="workflow-number">3</div>
                <div class="workflow-step-title">🎯 Match a Job</div>
                <div class="workflow-step-text">Optional: paste a job description and compare it with the resume.</div>
                <div class="workflow-next">Optional →</div>
            </div>
            <div class="workflow-step">
                <div class="workflow-number">4</div>
                <div class="workflow-step-title">🧠 Build Knowledge Base</div>
                <div class="workflow-step-text">Index the resume so the AI can search its content.</div>
                <div class="workflow-next">Required for chat →</div>
            </div>
            <div class="workflow-step">
                <div class="workflow-number">5</div>
                <div class="workflow-step-title">🤖 Ask Resume Agent</div>
                <div class="workflow-step-text">Chat naturally about experience, skills, projects and more.</div>
                <div class="workflow-next">Ask away →</div>
            </div>
        </div>
    </div>
    """,
    unsafe_allow_html=True,
)


# ============================================================
# TOP DASHBOARD: UPLOAD + ANALYSIS OVERVIEW
# ============================================================

left, right = st.columns([1.05, 1], gap="large")

with left:
    with st.container(border=True):
        render_section_title(
            "📄",
            "Upload Your Resume",
            "Upload PDF, DOC or DOCX to start the workflow",
        )

        st.markdown(
            """
            <div class="file-support">
                <span class="file-type">PDF</span>
                <span class="file-type">DOC</span>
                <span class="file-type">DOCX</span>
                <span style="color:#7b89a2;font-size:.72rem;padding-top:4px;">
                    Upload once, then click <b>Analyze Resume</b>.
                </span>
            </div>
            """,
            unsafe_allow_html=True,
        )

        uploaded_file = st.file_uploader(
            "Upload your resume",
            type=["pdf", "doc", "docx"],
            label_visibility="collapsed",
        )

        if uploaded_file is not None:
            is_new_file = uploaded_file.name != st.session_state.uploaded_file_name

            if is_new_file:
                try:
                    resume_text = prepare_uploaded_resume(uploaded_file)
                except Exception as exc:
                    st.error(f"❌ Could not read this resume: {exc}")
                    st.stop()

                # Keep a stable store id for compatibility with the existing RAG/agent backend.
                # Each rebuild overwrites the prior local store with the current uploaded resume.
                st.session_state.uploaded_file_name = uploaded_file.name
                st.session_state.resume_text = resume_text
                st.session_state.resume_id = RESUME_ID
                st.session_state.analysis = None
                st.session_state.job_result = None
                st.session_state.knowledge_base_ready = False
                st.session_state.knowledge_base_chunks = 0
                st.session_state.resume_conversation = []

            st.success(f"✅ {uploaded_file.name} uploaded successfully")
            st.caption("Step 2: Click **Analyze Resume** to generate the AI resume details. "
                       "Then optionally use **Analyze Match**, build the Knowledge Base, and chat with the Resume Agent.")

            b1, b2 = st.columns(2)
            with b1:
                if st.button("✨ Analyze Resume", use_container_width=True, help="Step 2: Generate ATS score, experience, strengths, missing skills and recommendations."):
                    with st.spinner("Analyzing resume with AI..."):
                        st.session_state.analysis = analyze_resume(
                            st.session_state.resume_text
                        )
                    st.toast("Resume analysis completed", icon="✅")

            with b2:
                if st.button("🧠 Build Knowledge Base", use_container_width=True, help="Step 4: Index your resume for AI search and Resume Agent chat."):
                    with st.spinner("Creating chunks and embeddings..."):
                        count = build_current_resume_store()
                    st.session_state.knowledge_base_ready = True
                    st.session_state.knowledge_base_chunks = count
                    st.toast(f"{count} chunks indexed", icon="🧠")

        else:
            st.info("Upload a PDF, DOC or DOCX resume to enable analysis, matching and AI search.")

with right:
    with st.container(border=True):
        render_section_title(
            "📊",
            "Analysis Overview",
            "Key insights from your uploaded resume",
        )

        analysis = st.session_state.analysis

        if analysis is None:
            m1, m2 = st.columns(2)
            with m1:
                render_metric_card("📈", "ATS Score", "--", "green")
            with m2:
                render_metric_card("📅", "Years of Experience", "--", "blue")

            m3, m4 = st.columns(2)
            with m3:
                render_metric_card("🧩", "Strengths Found", "--", "purple")
            with m4:
                render_metric_card("⚠️", "Missing Key Skills", "--", "red")

            st.caption("Run **Analyze Resume** to populate these insights.")
        else:
            top_a, top_b = st.columns([.9, 1.2])

            with top_a:
                render_score_ring(analysis.ats_score)

            with top_b:
                render_metric_card(
                    "📅",
                    "Years of Experience",
                    f"{analysis.years_of_experience} years",
                    "blue",
                )

            bottom_a, bottom_b = st.columns(2)
            with bottom_a:
                render_metric_card(
                    "🧩",
                    "Strengths Found",
                    str(len(safe_list(analysis.strengths))),
                    "purple",
                )

            with bottom_b:
                render_metric_card(
                    "⚠️",
                    "Missing Key Skills",
                    str(len(safe_list(analysis.missing_skills))),
                    "red",
                )


# ============================================================
# SECOND ROW: JD MATCH + QUICK INSIGHTS + AGENT TEASER
# ============================================================

st.markdown('<div id="job-description-match"></div>', unsafe_allow_html=True)

job_col, insight_col, agent_col = st.columns([1.05, .9, 1.15], gap="large")

with job_col:
    with st.container(border=True):
        render_section_title(
            "🎯",
            "Job Description Match",
            "See how well your resume fits a target role",
        )

        st.caption("Optional Step 3: Paste the job description below, then click **Analyze Match** to compare the role with your resume.")

        job_description = st.text_area(
            "Job description",
            value=st.session_state.job_description,
            height=185,
            placeholder="Paste the job description here...",
            label_visibility="collapsed",
        )
        st.session_state.job_description = job_description

        if st.button("🎯 Analyze Match", use_container_width=True, help="Optional Step 3: Compare your resume with a job description."):
            if not st.session_state.resume_text:
                st.warning("Please upload a resume first.")
            elif not job_description.strip():
                st.warning("Please paste a job description first.")
            else:
                with st.spinner("Comparing resume with the job description..."):
                    st.session_state.job_result = match_resume_to_job(
                        st.session_state.resume_text,
                        job_description,
                    )
                st.toast("Job match analysis completed", icon="🎯")

        if st.session_state.job_result is not None:
            result = st.session_state.job_result
            st.metric("Match Score", f"{result.match_score}/100")

with insight_col:
    with st.container(border=True):
        render_section_title(
            "💡",
            "Quick Insights",
            "What you can do with this app",
        )

        st.markdown(
            """
            <div class="insight-row">
                <div class="insight-icon">🧠</div>
                <div><div class="insight-title">Get ATS feedback</div><div class="insight-text">Understand how to improve your resume.</div></div>
            </div>
            <div class="insight-row">
                <div class="insight-icon">🎯</div>
                <div><div class="insight-title">Match job descriptions</div><div class="insight-text">Evaluate fit for target roles.</div></div>
            </div>
            <div class="insight-row">
                <div class="insight-icon">🧩</div>
                <div><div class="insight-title">Explore skills with AI</div><div class="insight-text">Identify strengths and development gaps.</div></div>
            </div>
            <div class="insight-row">
                <div class="insight-icon">💬</div>
                <div><div class="insight-title">Chat with your resume</div><div class="insight-text">Ask grounded questions and get instant answers.</div></div>
            </div>
            """,
            unsafe_allow_html=True,
        )

with agent_col:
    with st.container(border=True):
        render_section_title(
            "💬",
            "Chat with Your Resume Agent",
            "Ask anything about your experience, skills or projects",
        )

        st.markdown(
            """
            <div style="background:#f5f7fc;border-radius:13px;padding:14px 15px;margin-bottom:10px;">
                <div style="font-weight:800;color:#193266;margin-bottom:7px;">🤖 Hi! I'm your AI Resume Agent.</div>
                <div style="color:#65748f;font-size:.83rem;margin-bottom:7px;">You can ask questions like:</div>
                <div style="color:#43516d;font-size:.82rem;line-height:1.65;">
                    • What are my key skills?<br>
                    • How can I improve my resume?<br>
                    • Do I have experience with Python?<br>
                    • What AI projects are mentioned?<br>
                    • Suggest suitable roles from my resume.
                </div>
            </div>
            """,
            unsafe_allow_html=True,
        )

        if st.session_state.knowledge_base_ready:
            st.success("Resume Agent is ready.")
        else:
            st.caption("Build the Knowledge Base to activate the Resume Agent.")


# ============================================================
# RESUME ANALYSIS DETAILS
# ============================================================

st.markdown('<div id="resume-analysis"></div>', unsafe_allow_html=True)

if st.session_state.analysis is not None:
    analysis = st.session_state.analysis

    render_section_title(
        "📊",
        "Resume Analysis",
        "Detailed AI feedback from your uploaded resume",
    )

    c1, c2 = st.columns(2, gap="large")

    with c1:
        with st.container(border=True):
            render_list_card(
                "Strengths",
                "💪",
                safe_list(analysis.strengths),
            )

        with st.container(border=True):
            render_list_card(
                "Missing / Suggested Skills",
                "🧩",
                safe_list(analysis.missing_skills),
            )

    with c2:
        with st.container(border=True):
            render_list_card(
                "Areas to Improve",
                "⚠️",
                safe_list(analysis.weaknesses),
            )

        with st.container(border=True):
            render_list_card(
                "Suitable Roles",
                "🎯",
                safe_list(analysis.suitable_roles),
            )

    with st.expander("📋 Full AI Assessment", expanded=False):
        st.markdown("#### Recommended Skills")
        for item in safe_list(analysis.recommended_skills):
            st.markdown(f"- {item}")

        st.markdown("#### Resume Improvement Suggestions")
        for item in safe_list(analysis.improvement_suggestions):
            st.markdown(f"- {item}")

        st.markdown("#### Overall Assessment")
        st.write(analysis.overall_assessment)

        st.markdown("#### ⭐ Top Recommendation")
        st.info(analysis.top_recommendation)


# ============================================================
# JOB MATCH DETAILS
# ============================================================

if st.session_state.job_result is not None:
    result = st.session_state.job_result

    render_section_title(
        "🎯",
        "Job Match Analysis",
        "Detailed comparison between the resume and job description",
    )

    j1, j2, j3 = st.columns(3)
    with j1:
        render_metric_card("🎯", "Match Score", f"{result.match_score}/100", "green")
    with j2:
        render_metric_card("✅", "Matching Skills", str(len(safe_list(result.matching_skills))), "blue")
    with j3:
        render_metric_card("⚠️", "Missing Skills", str(len(safe_list(result.missing_skills))), "red")

    x1, x2 = st.columns(2, gap="large")
    with x1:
        with st.container(border=True):
            render_list_card("Matching Skills", "✅", safe_list(result.matching_skills))
            st.markdown("---")
            render_list_card("Matching Experience", "💼", safe_list(result.matching_experience))

    with x2:
        with st.container(border=True):
            render_list_card("Missing Skills", "❌", safe_list(result.missing_skills))
            st.markdown("---")
            render_list_card("Experience Gaps", "⚠️", safe_list(result.experience_gaps))

    with st.expander("📝 Recommendations & Overall Assessment"):
        for item in safe_list(result.recommendations):
            st.markdown(f"- {item}")
        st.markdown("#### Overall Assessment")
        st.write(result.overall_assessment)


# ============================================================
# KNOWLEDGE BASE + EXTRACTED PREVIEW
# ============================================================

st.markdown('<div id="resume-knowledge-base"></div>', unsafe_allow_html=True)

kb_col, preview_col = st.columns(2, gap="large")

with kb_col:
    with st.container(border=True):
        render_section_title(
            "🗄️",
            "Your Resume Knowledge Base",
            "Build a searchable knowledge base from your resume",
        )

        if st.session_state.knowledge_base_ready:
            st.markdown(
                f"""
                <div class="kb-ready">
                    <div class="kb-ready-title">✅ Knowledge base ready</div>
                    <div class="kb-ready-text">
                        Your resume is indexed and ready for AI search.
                        <b>{st.session_state.knowledge_base_chunks}</b> chunks are available.
                    </div>
                </div>
                """,
                unsafe_allow_html=True,
            )
        else:
            existing = load_vectors(st.session_state.resume_id)
            if existing and st.session_state.uploaded_file_name:
                st.session_state.knowledge_base_ready = True
                st.session_state.knowledge_base_chunks = len(existing)
                st.info(f"Existing vector store found with {len(existing)} chunks.")
            else:
                st.info("Upload a resume and build the Knowledge Base.")

        if st.session_state.resume_text:
            if st.button("🔄 Rebuild Knowledge Base", use_container_width=True, help="Re-index the currently uploaded resume."):
                with st.spinner("Rebuilding embeddings..."):
                    count = build_current_resume_store()
                st.session_state.knowledge_base_ready = True
                st.session_state.knowledge_base_chunks = count
                st.rerun()

with preview_col:
    with st.container(border=True):
        render_section_title(
            "📑",
            "Extracted Resume Preview",
            "Preview the text extracted from your resume",
        )

        if st.session_state.resume_text:
            st.caption("Readable text extracted from your uploaded resume. Use the full-text option below to inspect everything.")
            preview = st.session_state.resume_text[:1400]
            st.text_area(
                "Extracted Resume Preview",
                preview,
                height=240,
                label_visibility="collapsed",
            )

            with st.expander("👁 View Full Extracted Text"):
                st.text_area(
                    "Full extracted resume",
                    st.session_state.resume_text,
                    height=420,
                    label_visibility="collapsed",
                )
        else:
            st.info("Upload a resume to preview extracted text.")


# ============================================================
# RESUME AGENT
# ============================================================

st.markdown('<div id="resume-agent"></div>', unsafe_allow_html=True)

render_section_title(
    "🤖",
    "Resume Agent",
    "Ask grounded questions about the uploaded resume",
)

with st.container(border=True):
    if not st.session_state.knowledge_base_ready:
        st.info("🔒 Step 4 required: Click **Build Knowledge Base** above before chatting with the Resume Agent.")
    else:
        if not st.session_state.resume_conversation:
            st.markdown(
                """
                <div style="background:#f5f8ff;border:1px solid #e1e9f8;border-radius:13px;padding:14px 16px;margin-bottom:12px;">
                    <b style="color:#173369;">🤖 Resume Agent is ready.</b><br>
                    <span style="color:#71809a;font-size:.84rem;">
                        Ask about experience, projects, technologies, strengths, responsibilities or missing skills.
                    </span>
                </div>
                """,
                unsafe_allow_html=True,
            )

        for message in st.session_state.resume_conversation:
            with st.chat_message(message["role"]):
                st.markdown(message["content"])

                if message.get("results"):
                    with st.expander("🔍 Retrieved Resume Evidence"):
                        for item in message["results"]:
                            st.markdown(
                                f"**Chunk {item['chunk_id']}** · Similarity `{item['score']:.4f}`"
                            )
                            st.write(item["text"])
                            st.divider()

        question = st.chat_input(
            "Type your question about the resume..."
        )

        if question:
            previous_history = list(st.session_state.resume_conversation)

            st.session_state.resume_conversation.append(
                {"role": "user", "content": question}
            )

            with st.chat_message("user"):
                st.markdown(question)

            with st.chat_message("assistant"):
                with st.spinner("Searching resume evidence..."):
                    answer, results = answer_question(
                        question,
                        previous_history,
                        st.session_state.resume_id,
                    )

                st.markdown(answer)

                if results:
                    with st.expander("🔍 Retrieved Resume Evidence"):
                        for item in results:
                            st.markdown(
                                f"**Chunk {item['chunk_id']}** · Similarity `{item['score']:.4f}`"
                            )
                            st.write(item["text"])
                            st.divider()

            st.session_state.resume_conversation.append(
                {
                    "role": "assistant",
                    "content": answer,
                    "results": results,
                }
            )

        clear_col, spacer_col = st.columns([1, 3])
        with clear_col:
            if st.button("🗑 Clear Conversation", use_container_width=True):
                st.session_state.resume_conversation = []
                st.rerun()


# ============================================================
# ABOUT + FOOTER
# ============================================================

st.markdown('<div id="about"></div>', unsafe_allow_html=True)

footer_html = textwrap.dedent(f"""
<div class="app-footer">
    <div>
        <div><b>© 2026 AI Resume Reviewer.</b> All rights reserved.</div>
        <div style="margin-top:4px;">Built with ❤️ using OpenAI, Streamlit and Python.</div>
        <div class="footer-links" style="margin-top:7px;">
            <a href="{GITHUB_URL}" target="_blank">GitHub</a>
            <a href="{LIVE_URL}" target="_blank">Live Demo</a>
            <a href="{LINKEDIN_URL}" target="_blank">LinkedIn</a>
        </div>
    </div>

    <div class="footer-center">
        Automate Today,<br>
        <b>A Brighter Tomorrow ✨</b>
    </div>

    <div class="footer-name">
        Created by
        <strong>Sivani Sankar Mohapatra</strong>
    </div>
</div>
""").strip()

st.html(footer_html)
